from flask import Flask, render_template, request, redirect, url_for, session, send_file, flash
from functools import wraps
from datetime import datetime
import sqlite3
import os
import io

# Word generovanie
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# PDF generovanie
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image as RLImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Jednotný vzhľad tabuliek v správe o vykurovaní
TABLE_SHADE_HEX = 'E8F4EC'                 # jemné zelené podfarbenie (prvý stĺpec + hlavička)
TABLE_SHADE_RL = colors.HexColor('#e8f4ec')
FIRST_COL_CM = 6.0                          # jednotná šírka prvého stĺpca vo všetkých tabuľkách (cm)
TABLE_TOTAL_CM = 17.0

# Font s podporou slovenskej diakritiky (Liberation Serif - metricky kompatibilný s Times New Roman)
FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fonts')
PDF_FONT = 'Times-Roman'
PDF_FONT_BOLD = 'Times-Bold'
try:
    pdfmetrics.registerFont(TTFont('PDFSerif', os.path.join(FONT_DIR, 'LiberationSerif-Regular.ttf')))
    pdfmetrics.registerFont(TTFont('PDFSerif-Bold', os.path.join(FONT_DIR, 'LiberationSerif-Bold.ttf')))
    pdfmetrics.registerFontFamily('PDFSerif', normal='PDFSerif', bold='PDFSerif-Bold')
    PDF_FONT = 'PDFSerif'
    PDF_FONT_BOLD = 'PDFSerif-Bold'
except Exception:
    pass  # ak font chýba, použije sa vstavaný Times-Roman (bez slovenskej diakritiky)

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'zmen-toto-heslo-pred-nasadenim')

# ============================================================
# PRIHLASENIE - zmen USERNAME a PASSWORD
# ============================================================
USERNAME = os.environ.get('APP_USERNAME', 'admin')
PASSWORD = os.environ.get('APP_PASSWORD', 'heslo123')

# ============================================================
# DATABAZA
# ============================================================
DB_PATH = 'audity.db'

def init_db():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS audity (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        typ TEXT NOT NULL,
        datum TEXT NOT NULL,
        nazov_budovy TEXT,
        adresa TEXT,
        vlastnik TEXT,
        udaje TEXT,
        vysledky TEXT
    )''')
    conn.commit()
    conn.close()

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

# ============================================================
# VZORCE - ENERGETICKY AUDIT
# ============================================================
def vypocitaj_audit(data):
    try:
        plocha = float(data.get('plocha', 0))
        spotreba_teplo = float(data.get('spotreba_teplo', 0))
        spotreba_elektrina = float(data.get('spotreba_elektrina', 0))
        spotreba_celkova = spotreba_teplo + spotreba_elektrina

        # Merná spotreba energie
        if plocha > 0:
            merna_spotreba = spotreba_celkova / plocha
        else:
            merna_spotreba = 0

        # Energetická trieda
        if merna_spotreba < 30:
            trieda = "A0"
        elif merna_spotreba < 60:
            trieda = "A1"
        elif merna_spotreba < 100:
            trieda = "B"
        elif merna_spotreba < 150:
            trieda = "C"
        elif merna_spotreba < 200:
            trieda = "D"
        elif merna_spotreba < 250:
            trieda = "E"
        else:
            trieda = "F"

        return {
            'merna_spotreba': round(merna_spotreba, 2),
            'spotreba_celkova': round(spotreba_celkova, 2),
            'trieda': trieda
        }
    except:
        return {'merna_spotreba': 0, 'spotreba_celkova': 0, 'trieda': 'N/A'}

# ============================================================
# VZORCE - ENERGETICKY CERTIFIKAT
# ============================================================
def vypocitaj_certifikat(data):
    try:
        plocha = float(data.get('plocha', 0))
        potreba_tepla = float(data.get('potreba_tepla', 0))
        potreba_chlad = float(data.get('potreba_chlad', 0))
        potreba_svietenie = float(data.get('potreba_svietenie', 0))

        celkova_potreba = potreba_tepla + potreba_chlad + potreba_svietenie

        if plocha > 0:
            merna_potreba = celkova_potreba / plocha
        else:
            merna_potreba = 0

        # Energetická trieda certifikátu (podľa vyhlášky 364/2012)
        if merna_potreba < 32:
            trieda = "A0"
        elif merna_potreba < 54:
            trieda = "A1"
        elif merna_potreba < 95:
            trieda = "B"
        elif merna_potreba < 150:
            trieda = "C"
        elif merna_potreba < 210:
            trieda = "D"
        elif merna_potreba < 270:
            trieda = "E"
        else:
            trieda = "F"

        return {
            'celkova_potreba': round(celkova_potreba, 2),
            'merna_potreba': round(merna_potreba, 2),
            'trieda': trieda
        }
    except:
        return {'celkova_potreba': 0, 'merna_potreba': 0, 'trieda': 'N/A'}

# ============================================================
# VZORCE - KONTROLA VYKUROVANIA
# podľa Zákona č. 314/2012 Z.z. a Vyhl. č. 422/2012 Z.z.
# (pravidelná kontrola kotla a vykurovacieho systému)
# ============================================================

# Siegertove koeficienty komínovej straty pre jednotlivé druhy paliva
KOEFICIENTY_PALIVA = {
    'Zemný plyn':                {'A2': 0.65, 'B': 0.008, 'vyhrevnost': 9.86},
    'Vykurovací olej':           {'A2': 0.68, 'B': 0.007, 'vyhrevnost': 11.86},
    'Tuhé palivo - čierne uhlie':{'A2': 0.60, 'B': 0.010, 'vyhrevnost': 7.00},
    'Tuhé palivo - drevo/biomasa':{'A2': 0.65, 'B': 0.025, 'vyhrevnost': 4.00},
}

STRATA_SALANIM = 0.8  # % - typická strata sálaním kotla do okolia (odhad)

def komin_strata(t_spalin, t_vzduch, o2, palivo):
    """Komínová strata (Siegertova metóda) v %."""
    k = KOEFICIENTY_PALIVA.get(palivo, KOEFICIENTY_PALIVA['Zemný plyn'])
    if o2 >= 20.9 or t_spalin <= t_vzduch:
        return 0.0
    return (t_spalin - t_vzduch) * (k['A2'] / (20.9 - o2) + k['B'])

def _f(data, key, default=0):
    try:
        v = data.get(key)
        return float(v) if v not in (None, '') else default
    except (TypeError, ValueError):
        return default

def vypocitaj_vykurovanie(data):
    try:
        min_ucinnost = _f(data, 'min_ucinnost', 96)

        # --- 1) Meranie účinnosti kotlov pri 50 % a 100 % zaťažení (komínová strata) ---
        kotly = []
        for pref in ['k1', 'k2']:
            if not data.get(f'{pref}_menovity_vykon'):
                continue
            palivo = data.get('druh_paliva') or 'Zemný plyn'
            garantovana = _f(data, f'{pref}_garantovana_ucinnost')

            merania = []
            for bod, label in [('50', '50 %'), ('100', '100 %')]:
                t_spalin = _f(data, f'{pref}_t_spalin_{bod}')
                t_vzduch = _f(data, f'{pref}_t_vzduch_{bod}')
                o2 = _f(data, f'{pref}_o2_{bod}')
                if t_spalin == 0 and t_vzduch == 0 and o2 == 0:
                    continue
                strata_citelnym_teplom = round(komin_strata(t_spalin, t_vzduch, o2, palivo), 3)
                strata_salanim = _f(data, f'{pref}_strata_salanim_{bod}', STRATA_SALANIM)
                strata_horlavinou = _f(data, f'{pref}_strata_horlavinou_{bod}', 0)
                ucinnost_bod = round(100 - strata_citelnym_teplom - strata_salanim - strata_horlavinou, 2)
                prebytok_vzduchu = round(21 / (21 - o2), 2) if 0 < o2 < 21 else 0
                merania.append({
                    'zatazenie': label, 't_spalin': t_spalin, 't_vzduch': t_vzduch, 'o2': o2,
                    'co': data.get(f'{pref}_co_{bod}', ''), 'co2': data.get(f'{pref}_co2_{bod}', ''),
                    'so2': data.get(f'{pref}_so2_{bod}', ''), 'no': data.get(f'{pref}_no_{bod}', ''),
                    'no2': data.get(f'{pref}_no2_{bod}', ''),
                    'prebytok_vzduchu': prebytok_vzduchu,
                    'strata_salanim': strata_salanim, 'strata_horlavinou': strata_horlavinou,
                    'strata_citelnym_teplom': strata_citelnym_teplom,
                    'komin_strata': round(strata_citelnym_teplom, 2),
                    'ucinnost': ucinnost_bod,
                })

            if merania:
                priemer = round(sum(m['ucinnost'] for m in merania) / len(merania), 2)
            else:
                priemer = 0
            stav = 'vyhovuje' if priemer >= min_ucinnost else 'nevyhovuje'

            kotly.append({
                'oznacenie': pref.upper(),
                'vyrobca': data.get(f'{pref}_vyrobca', ''),
                'typ': data.get(f'{pref}_typ', ''),
                'vyrobne_cislo': data.get(f'{pref}_vyrobne_cislo', ''),
                'rok_vyroby': data.get(f'{pref}_rok_vyroby', ''),
                'prevadzkovy_stav': data.get(f'{pref}_prevadzkovy_stav') or 'v prevádzke',
                'menovity_vykon': data.get(f'{pref}_menovity_vykon', ''),
                'max_vykon_kondenzacny': data.get(f'{pref}_max_vykon_kondenzacny', ''),
                'max_vykon': data.get(f'{pref}_max_vykon', ''),
                'max_prikon': data.get(f'{pref}_max_prikon', ''),
                'min_vykon': data.get(f'{pref}_min_vykon', ''),
                'min_prikon': data.get(f'{pref}_min_prikon', ''),
                'kondenzacny': data.get(f'{pref}_kondenzacny') or 'kondenzačný',
                'oznacenie_ce': data.get(f'{pref}_oznacenie_ce', ''),
                'vyrobca_horaka': data.get(f'{pref}_vyrobca_horaka', '') or '-',
                'typ_horaka': data.get(f'{pref}_typ_horaka', '') or '-',
                'vyrobne_cislo_horaka': data.get(f'{pref}_vyrobne_cislo_horaka', '') or '-',
                'rok_vyroby_horaka': data.get(f'{pref}_rok_vyroby_horaka', '') or '-',
                'overenie_min': data.get(f'{pref}_overenie_min', ''),
                'overenie_max': data.get(f'{pref}_overenie_max', ''),
                'typ_paliva': palivo,
                'merania': merania,
                'priemerna_ucinnost': priemer,
                'garantovana_ucinnost': garantovana,
                'stav': stav,
            })

        # --- 2) Priama metóda — až 3 roky (rovnako ako v protokole) ---
        roky = []
        for i in range(1, 4):
            rok = data.get(f'rok{i}')
            odber = _f(data, f'rok{i}_odber_paliva')
            if not rok or odber == 0:
                continue
            vyhrevnost = _f(data, f'rok{i}_vyhrevnost')
            spalovacie_teplo = _f(data, f'rok{i}_spalovacie_teplo')
            vyrobene_teplo = _f(data, f'rok{i}_vyrobene_teplo')
            spotreba_vykurovanie = _f(data, f'rok{i}_spotreba_vykurovanie')
            spotreba_tuv = _f(data, f'rok{i}_spotreba_tuv')
            spotreba_vody_tuv = _f(data, f'rok{i}_spotreba_vody_tuv')

            pomer_vyhrevnosti = round(vyhrevnost / spalovacie_teplo, 8) if spalovacie_teplo > 0 else 0
            teplo_v_palive = round(odber * vyhrevnost, 2)
            ucinnost_priama = round((vyrobene_teplo / teplo_v_palive) * 100, 2) if teplo_v_palive > 0 else 0
            merna_spotreba_tuv = round(spotreba_tuv / spotreba_vody_tuv, 2) if spotreba_vody_tuv > 0 else 0
            celkova_spotreba = spotreba_vykurovanie + spotreba_tuv
            podiel_tuv = round((spotreba_tuv / celkova_spotreba) * 100, 2) if celkova_spotreba > 0 else 0

            # Spotreba paliva rozpočítaná na vykurovanie/TÚV podľa dosiahnutej účinnosti
            ucinnost_podiel = ucinnost_priama / 100 if ucinnost_priama > 0 else 1
            spotreba_paliva_vykurovanie = round(spotreba_vykurovanie / ucinnost_podiel, 2)
            spotreba_paliva_tuv = round(spotreba_tuv / ucinnost_podiel, 2)

            # Posúdenie potrebného výkonu (STN 06 0210 — stupeň-deňová metóda)
            ti = _f(data, f'rok{i}_ti', 20)
            te = _f(data, f'rok{i}_te', -15)
            tepr = _f(data, f'rok{i}_tepr')
            d = _f(data, f'rok{i}_pocet_dni', 200)
            if d > 0 and (ti - tepr) != 0:
                potrebny_vykon = round(spotreba_vykurovanie * (ti - te) / (24 * d * (ti - tepr)), 2)
            else:
                potrebny_vykon = 0
            pomer_ti = round((ti - tepr) / (ti - te), 3) if (ti - te) != 0 else 0

            roky.append({
                'rok': rok, 'odber_paliva': odber, 'vyhrevnost': vyhrevnost,
                'spalovacie_teplo': spalovacie_teplo, 'pomer_vyhrevnosti': pomer_vyhrevnosti,
                'teplo_v_palive': teplo_v_palive, 'vyrobene_teplo': vyrobene_teplo,
                'ucinnost_priama': ucinnost_priama,
                'spotreba_vykurovanie': spotreba_vykurovanie, 'spotreba_tuv': spotreba_tuv,
                'merna_spotreba_tuv': merna_spotreba_tuv, 'podiel_tuv': podiel_tuv,
                'spotreba_paliva_vykurovanie': spotreba_paliva_vykurovanie,
                'spotreba_paliva_tuv': spotreba_paliva_tuv,
                'ti': ti, 'te': te, 'tepr': tepr, 'pocet_dni': d,
                'gj_rok': round(spotreba_vykurovanie * 3.6 / 1000, 3),
                'pomer_ti': pomer_ti,
                'potrebny_vykon': potrebny_vykon,
            })

        priemerna_ucinnost_priama = round(sum(r['ucinnost_priama'] for r in roky) / len(roky), 2) if roky else 0

        instalovany_vykon = round(sum(_f(data, f'{p}_menovity_vykon') for p in ['k1', 'k2']), 2)
        posledny_potrebny_vykon = roky[-1]['potrebny_vykon'] if roky else 0
        vykon_rezerva = round(instalovany_vykon - posledny_potrebny_vykon, 2)

        # --- Celkové vyhodnotenie ---
        if kotly:
            priemerna_ucinnost = round(sum(k['priemerna_ucinnost'] for k in kotly) / len(kotly), 2)
            celkovy_stav = 'vyhovuje' if all(k['stav'] == 'vyhovuje' for k in kotly) else 'nevyhovuje'
        else:
            priemerna_ucinnost = 0
            celkovy_stav = 'N/A'

        # --- Termín nasledujúcej kontroly ---
        interval = int(data.get('interval_kontroly') or 4)
        nasledujuca_kontrola = ''
        datum_kontroly = data.get('datum_kontroly', '')
        try:
            mesiac, rok_k = datum_kontroly.split('/')
            nasledujuca_kontrola = f"{mesiac}/{int(rok_k) + interval}"
        except Exception:
            nasledujuca_kontrola = ''

        return {
            'kotly': kotly,
            'min_ucinnost': min_ucinnost,
            'priemerna_ucinnost': priemerna_ucinnost,
            'celkovy_stav': celkovy_stav,
            'roky': roky,
            'priemerna_ucinnost_priama': priemerna_ucinnost_priama,
            'instalovany_vykon': instalovany_vykon,
            'potrebny_vykon': posledny_potrebny_vykon,
            'vykon_rezerva': vykon_rezerva,
            'interval_kontroly': interval,
            'nasledujuca_kontrola': nasledujuca_kontrola,
        }
    except Exception:
        return {
            'kotly': [], 'min_ucinnost': 96, 'priemerna_ucinnost': 0,
            'celkovy_stav': 'N/A', 'roky': [], 'priemerna_ucinnost_priama': 0,
            'instalovany_vykon': 0, 'potrebny_vykon': 0, 'vykon_rezerva': 0,
            'interval_kontroly': 4, 'nasledujuca_kontrola': '',
        }

# ============================================================
# ZOZNAM TABULIEK - jednotný zdroj poradia a popiskov (Word aj PDF)
# ============================================================
def _tabulky_zoznam(kotly, roky, ma_oze, ma_tc):
    zoznam = [
        "Identifikačné údaje kontrolovaného kotla (spoločné údaje)",
        "Identifikačné údaje kontrolovaného kotla (parametre kotlov)",
        "Identifikačné údaje kontrolovaného kotla (regulácia a médiá)",
    ]
    if ma_oze:
        zoznam.append("Identifikačné údaje ostatných zariadení na výrobu tepla – OZE")
    if ma_tc:
        zoznam.append("Identifikačné údaje ostatných zariadení na výrobu tepla – tepelné čerpadlo")
    zoznam += [
        "Prevádzková dokumentácia kotlov",
        "Vizuálna kontrola a zhodnotenie kotla",
        "Zhodnotenie údržby a kontrola dokladov o údržbe a opravách kotla",
        "Kontrola funkčnosti kotla",
        "Kontrola úplnosti a aktuálnosti dokumentácie vnútorných rozvodov",
        "Prehliadka vnútorných rozvodov tepla a teplej vody",
        "Zhodnotenie údržby rozvodov tepla a teplej vody",
        "Porovnanie skutočného využívania budovy s projektovaným využívaním",
        "Porovnanie skutočného využívania rozvodov tepla s projektovaným využívaním",
    ]
    for k in kotly:
        if k['merania']:
            zoznam.append(f"Namerané a vypočítané parametre – kotol {k['oznacenie']}")
    if roky:
        zoznam.append("Vyhodnotenie účinnosti výroby tepla priamou metódou")
        zoznam.append("Zhodnotenie priemerného tepelného výkonu")
    return zoznam

# ============================================================
# GENEROVANIE WORD - pomocné funkcie pre tabuľky
# ============================================================
def _cell(cell, text, bold=False, size=9, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run('' if text is None else str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'

def _shade_cell(cell, hex_color=TABLE_SHADE_HEX):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def _zero_table_indent(table):
    """Zarovná ľavý okraj tabuľky presne s ľavým okrajom bežného textu (odstráni tblInd)."""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    existing = tblPr.find(qn('w:tblInd'))
    if existing is not None:
        tblPr.remove(existing)
    tblInd = OxmlElement('w:tblInd')
    tblInd.set(qn('w:w'), '0')
    tblInd.set(qn('w:type'), 'dxa')
    tblPr.append(tblInd)

def _docx_table_caption(doc, counter, caption):
    counter[0] += 1
    p = doc.add_paragraph(f"Tabuľka č. {counter[0]} – {caption}" if caption else f"Tabuľka č. {counter[0]}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
    return counter[0]

def _docx_table(doc, header_row, data_rows, counter=None, caption=None):
    if counter is not None:
        _docx_table_caption(doc, counter, caption)
    table = doc.add_table(rows=1, cols=len(header_row))
    table.style = 'Table Grid'
    table.autofit = False
    _zero_table_indent(table)
    for i, h in enumerate(header_row):
        _cell(table.rows[0].cells[i], h, bold=True)
        _shade_cell(table.rows[0].cells[i])
    table.rows[0].cells[0].width = Cm(FIRST_COL_CM)
    for row in data_rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            _cell(cells[i], val)
        _shade_cell(cells[0])
        cells[0].width = Cm(FIRST_COL_CM)
    doc.add_paragraph("")
    return table

def _docx_section_heading(doc, text):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

def _docx_kotly_identifikacia(doc, kotly, counter=None, caption=None):
    header = ["Označenie kotla"] + [k['oznacenie'] for k in kotly]
    riadky = [
        ("Prevádzkový stav", 'prevadzkovy_stav'), ("Výrobca kotla", 'vyrobca'), ("Typ kotla", 'typ'),
        ("Výrobné číslo kotla", 'vyrobne_cislo'), ("Rok výroby kotla", 'rok_vyroby'),
        ("Menovitý výkon", 'menovity_vykon'), ("Maximálny výkon 50/30°C (kondenzačný)", 'max_vykon_kondenzacny'),
        ("Maximálny výkon 80/60 °C", 'max_vykon'),
        ("Maximálny príkon", 'max_prikon'), ("Minimálny výkon 80/60 °C", 'min_vykon'),
        ("Minimálny príkon", 'min_prikon'), ("Kondenzačný/nekondenzačný", 'kondenzacny'),
        ("Označenie CE", 'oznacenie_ce'),
        ("Výrobca horáka, ak je kotol vybavený horákom dodatočne", 'vyrobca_horaka'),
        ("Typ horáka", 'typ_horaka'), ("Výrobné číslo horáka", 'vyrobne_cislo_horaka'),
        ("Rok výroby horáka", 'rok_vyroby_horaka'),
    ]
    data_rows = []
    for label, key in riadky:
        data_rows.append([label] + [k.get(key, '') for k in kotly])
    _docx_table(doc, header, data_rows, counter=counter, caption=caption)

def _add_bottom_border(paragraph):
    """Súvislé podčiarknutie cez celú šírku strany (hlavička dokumentu)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)

# ============================================================
# GENEROVANIE WORD
# ============================================================
def generuj_word(typ, data, vysledky):
    doc = Document()

    # Štýl nadpisu
    nazov = {
        'audit': 'ENERGETICKÝ AUDIT',
        'certifikat': 'ENERGETICKÝ CERTIFIKÁT',
        'vykurovanie': 'Správa z vykonanej pravidelnej kontroly vykurovacieho systému'
    }.get(typ, 'SPRÁVA')

    h = doc.add_heading(nazov, 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)

    if typ == 'vykurovanie':
        kotly = vysledky.get('kotly', [])
        roky = vysledky.get('roky', [])
        ma_oze = any(data.get(k) for k in ['oze_druh_energie', 'oze_druh_zariadenia', 'oze_vyrobca'])
        ma_tc = any(data.get(k) for k in ['tc_vyrobca', 'tc_typ', 'tc_prevadzkovy_stav'])
        tc = [0]  # zdieľaný počítadlo tabuliek

        sub = doc.add_paragraph("podľa Zákona č. 314/2012 Z.z.")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in sub.runs:
            run.font.name = 'Times New Roman'
        doc.add_paragraph("")

        for label, key, extra in [
            ("Objekt", 'nazov_budovy', None), ("Lokalita", 'adresa', None),
            ("Majiteľ", 'vlastnik', None), ("Adresa vlastníka", 'adresa_vlastnika', None),
            ("Správca", 'spravca', None), ("Prevádzkovateľ", 'prevadzkovatel', None),
        ]:
            p = doc.add_paragraph()
            r1 = p.add_run(f"{label}: ")
            r1.bold = True
            r1.font.name = 'Times New Roman'
            r2 = p.add_run(str(data.get(key, '')))
            r2.font.name = 'Times New Roman'
        doc.add_paragraph("")

        fotografia = data.get('fotografia_cesta')
        if fotografia and os.path.isfile(fotografia):
            doc.add_picture(fotografia, width=Inches(4))
            pic_p = doc.paragraphs[-1]
            pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("")

        for label, key in [
            ("Vypracoval", 'vypracoval_firma'), ("Kontakt", 'vypracoval_kontakt'),
            ("Dátum kontroly", 'datum_kontroly'), ("Poradové číslo", 'poradove_cislo'),
        ]:
            p = doc.add_paragraph()
            r1 = p.add_run(f"{label}: ")
            r1.bold = True
            r1.font.name = 'Times New Roman'
            r2 = p.add_run(str(data.get(key, '')))
            r2.font.name = 'Times New Roman'

        doc.add_page_break()

        # --- OBSAH (strana 2 - bez hlavičky/päty) ---
        _docx_section_heading(doc, "Obsah")
        for polozka in [
            "1. Kontrola kotla",
            "    1.1 Identifikačné údaje kontrolovaného kotla",
            "    1.2 Identifikačné údaje ostatných zariadení na výrobu tepla",
            "2. Vizuálna kontrola a zhodnotenie funkčnosti a údržby kotla",
            "    2.1 Prevádzková dokumentácia kotlov a povinnosti z nej vyplývajúce",
            "    2.2 Vizuálna kontrola a zhodnotenie kotla",
            "    2.3 Zhodnotenie údržby a kontrola dokladov o údržbe a opravách",
            "    2.4 Kontrola funkčnosti kotla",
            "3. Rozšírená kontrola vykurovacieho systému",
            "    3.1 Kontrola úplnosti a aktuálnosti dokumentácie vnútorných rozvodov tepla a teplej vody",
            "    3.2 Prehliadka vnútorných rozvodov tepla a teplej vody",
            "    3.3 Zhodnotenie údržby a kontrola dokladov o údržbe a opravách rozvodov tepla a teplej vody",
            "    3.4 Porovnanie skutočného využívania budovy s projektovaným využívaním",
            "    3.5 Porovnanie skutočného využívania rozvodov tepla s projektovaným využívaním",
            "4. Meranie účinnosti kotla – komínová strata",
            "5. Vyhodnotenie účinnosti výroby tepla priamou metódou",
            "6. Posúdenie výkonu kotla vzhľadom na potrebu tepla v budove",
            "7. Vyhodnotenie kontroly kotla a návrhy na opatrenia",
            "8. Vyhodnotenie rozšírenej kontroly vykurovacieho systému a návrh opatrení",
        ]:
            po = doc.add_paragraph(polozka)
            for run in po.runs:
                run.font.name = 'Times New Roman'

        doc.add_paragraph("")
        zoznam_nadpis = doc.add_paragraph("Zoznam tabuliek")
        zoznam_nadpis.runs[0].bold = True
        zoznam_nadpis.runs[0].font.name = 'Times New Roman'
        for i, caption in enumerate(_tabulky_zoznam(kotly, roky, ma_oze, ma_tc), start=1):
            zp = doc.add_paragraph(f"Tabuľka č. {i} – {caption}")
            for run in zp.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9)

        # --- Nová sekcia dokumentu (od strany 3): hlavička + päta ---
        section2 = doc.add_section(WD_SECTION.NEW_PAGE)
        section2.header.is_linked_to_previous = False
        section2.footer.is_linked_to_previous = False

        hp = section2.header.paragraphs[0]
        hp.text = f"{data.get('nazov_budovy', '')}, {data.get('adresa', '')}"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in hp.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9)
        _add_bottom_border(hp)

        fp = section2.footer.paragraphs[0]
        fp.text = data.get('vypracoval_firma', '')
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in fp.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9)

        # --- 1. KONTROLA KOTLA ---
        _docx_section_heading(doc, "1. Kontrola kotla")
        doc.add_paragraph("Kontrola kotla bola vykonávaná podľa Vyhl. č. 422/2012 Z.z. §2 v nasledovnom členení:")
        doc.add_paragraph("1.1 Identifikačné údaje kontrolovaného kotla").runs[0].bold = True
        _docx_table(doc, ["", ""], [
            ("Vlastník", data.get('vlastnik', '')), ("Adresa vlastníka", data.get('adresa_vlastnika', '')),
            ("Adresa budovy, v ktorej je kotol umiestnený", data.get('adresa', '')),
            ("Správca", data.get('spravca', '')), ("Prevádzkovateľ", data.get('prevadzkovatel', '')),
            ("Typ paliva", data.get('typ_paliva_kategoria', '')), ("Druh paliva", data.get('druh_paliva', '')),
            ("Spôsob dávkovania paliva", data.get('sposob_davkovania', '')),
        ], counter=tc, caption="Identifikačné údaje kontrolovaného kotla (spoločné údaje)")
        if kotly:
            _docx_kotly_identifikacia(doc, kotly, counter=tc, caption="Identifikačné údaje kontrolovaného kotla (parametre kotlov)")
        else:
            doc.add_paragraph("Nebol zadaný žiadny kotol.")
        _docx_table(doc, ["", ""], [
            ("Typ výkonovej regulácie", data.get('typ_regulacie', '')),
            ("Spôsob odvodu spalín", data.get('sposob_odvodu_spalin', '')),
            ("Spôsob prívodu vzduchu", data.get('sposob_privodu_vzduchu', '')),
            ("Teplonosné médium", data.get('teplonosne_medium', '')),
            ("Spôsob využitia kotla", data.get('sposob_vyuzitia', '')),
        ], counter=tc, caption="Identifikačné údaje kontrolovaného kotla (regulácia a médiá)")

        doc.add_paragraph("1.2 Identifikačné údaje ostatných zariadení na výrobu tepla").runs[0].bold = True
        doc.add_paragraph("V budove je inštalované zariadenie na využívanie OZE." if ma_oze else "V budove nie je inštalované zariadenie na využívanie OZE.")
        if ma_oze:
            _docx_table(doc, ["", ""], [
                ("Druh využívanej energie", data.get('oze_druh_energie', '')),
                ("Druh zariadenia", data.get('oze_druh_zariadenia', '')),
                ("Výrobca", data.get('oze_vyrobca', '')),
                ("Apertúrna plocha kolektora", data.get('oze_apertura', '')),
                ("Počet kusov", data.get('oze_pocet_kusov', '')),
                ("Celkový inštalovaný výkon", data.get('oze_celkovy_vykon', '')),
            ], counter=tc, caption="Identifikačné údaje ostatných zariadení na výrobu tepla – OZE")
        if ma_tc:
            doc.add_paragraph("Tepelné čerpadlo")
            _docx_table(doc, ["", ""], [
                ("Prevádzkový stav", data.get('tc_prevadzkovy_stav', '')), ("Výrobca", data.get('tc_vyrobca', '')),
                ("Typ", data.get('tc_typ', '')), ("Menovitý výkon", data.get('tc_menovity_vykon', '')),
                ("Menovitý príkon", data.get('tc_menovity_prikon', '')), ("COP", data.get('tc_cop', '')),
            ], counter=tc, caption="Identifikačné údaje ostatných zariadení na výrobu tepla – tepelné čerpadlo")
        if data.get('poznamky_zariadenia'):
            doc.add_paragraph(f"Poznámky: {data.get('poznamky_zariadenia')}")

        # --- 2. VIZUÁLNA KONTROLA ---
        _docx_section_heading(doc, "2. Vizuálna kontrola a zhodnotenie funkčnosti a údržby kotla")
        doc.add_paragraph("2.1 Prevádzková dokumentácia kotlov a povinnosti z nej vyplývajúce").runs[0].bold = True
        _docx_table(doc, ["Dokumentácia kotla", ""], [
            ("Projektová dokumentácia kotla", data.get('dok_projektova', '')),
            ("Prevádzkový predpis výr. kotla", data.get('dok_predpis', '')),
            ("Dokumentácia prevádzky a údržby", data.get('dok_udrzba', '')),
            ("Správa z predchádzajúcej kontroly", data.get('dok_predchadzajuca_sprava', '')),
            ("Zaškolenie obsluhy", data.get('dok_zaskolenie', '')),
            ("Odborné prehliadky a revízie", data.get('dok_revizie', '')),
        ], counter=tc, caption="Prevádzková dokumentácia kotlov")
        doc.add_paragraph("2.2 Vizuálna kontrola a zhodnotenie kotla").runs[0].bold = True
        _docx_table(doc, ["", ""], [
            ("Únik paliva", data.get('unik_paliva', '')), ("Únik teplonosnej látky", data.get('unik_teplonosnej_latky', '')),
            ("Vonkajší stav kotla", data.get('vonkajsi_stav_kotla', '')),
            ("Tepelná izolácia, oplechovanie, netesnosti spalinovodu", data.get('izolacia_oplechovanie', '')),
            ("Znečistenie spaľovacej komory a teplovýmenných plôch", data.get('znecistenie', '')),
            ("Znečistenie horákov", data.get('znecistenie_horakov', '')),
            ("Funkčnosť armatúr a stav ostatných častí vyžadujúcich údržbu", data.get('armatury', '')),
            ("Kvalita teplonosnej látky, čistota obehovej vody", data.get('kvalita_vody', '')),
            ("Správnosť údajov meracích prístrojov", data.get('meracie_pristroje', '')),
            ("Systém riadenia kotla", data.get('system_riadenia', '')),
            ("Čistota a poriadok kotolne", data.get('cistota_kotolne', '')),
        ], counter=tc, caption="Vizuálna kontrola a zhodnotenie kotla")
        doc.add_paragraph("2.3 Zhodnotenie údržby a kontrola dokladov o údržbe a opravách").runs[0].bold = True
        _docx_table(doc, ["", ""], [
            ("Zhodnotenie údržby a zjavných stôp po údržbárskych prácach", data.get('udrzba_stopy', '')),
            ("Kontrola dokladov o údržbe a opravách", data.get('doklady_udrzba', '')),
        ], counter=tc, caption="Zhodnotenie údržby a kontrola dokladov o údržbe a opravách kotla")
        doc.add_paragraph("2.4 Kontrola funkčnosti kotla").runs[0].bold = True
        funkcnost_rows = [("Skúška funkcie kotlov v prevádzke", data.get('funkcnost_skuska', ''))]
        for k in kotly:
            if k['overenie_min'] or k['overenie_max']:
                funkcnost_rows.append((f"Overenie výkonu kotla {k['oznacenie']}",
                                        f"min = {k['overenie_min']} m³/hod   max = {k['overenie_max']} m³/hod"))
        funkcnost_rows.append(("Regulácia výkonu", data.get('regulacia_vykonu_text', '')))
        _docx_table(doc, ["", ""], funkcnost_rows, counter=tc, caption="Kontrola funkčnosti kotla")

        # --- 3. ROZŠÍRENÁ KONTROLA ---
        _docx_section_heading(doc, "3. Rozšírená kontrola vykurovacieho systému")
        doc.add_paragraph("Rozšírená kontrola vykurovacieho systému bola vykonávaná podľa Vyhl. č. 422/2012 Z.z. §3 v nasledovnom členení:")
        doc.add_paragraph("3.1 Kontrola úplnosti a aktuálnosti dokumentácie vnútorných rozvodov tepla a teplej vody").runs[0].bold = True
        _docx_table(doc, ["", ""], [
            ("Úplnosť PD ÚK a TÚV", data.get('roz_pd_uk_tuv', '')), ("Úplnosť PD zmien a rekonštrukcie", data.get('roz_pd_zmien', '')),
            ("Úplnosť prevádz. predpisov zariadení", data.get('roz_predpisy', '')), ("Úplnosť MPP", data.get('roz_mpp', '')),
            ("Vedenie prevádzkového denníka", data.get('roz_dennik', '')), ("Správy o údržbe a opravách", data.get('roz_spravy_udrzba', '')),
            ("Správa z predchádzajúcej kontroly", data.get('dok_predchadzajuca_sprava', '')),
        ], counter=tc, caption="Kontrola úplnosti a aktuálnosti dokumentácie vnútorných rozvodov")
        doc.add_paragraph("3.2 Prehliadka vnútorných rozvodov tepla a teplej vody").runs[0].bold = True
        _docx_table(doc, ["", ""], [
            ("Hlavné komponenty rozv. tepla, prvky merania a riadenia", data.get('roz_komponenty', '')),
            ("Ovládacie prvky systému regulácie", data.get('roz_ovladacie_prvky', '')),
            ("Vykurovacie telesá", data.get('vykurovacie_telesa', '')),
            ("Tepelná izolácia rozvodov tepla", data.get('tepelna_izolacia_rozvodov', '')),
            ("Čistota obehovej vody", data.get('cistota_obehovej_vody', '')),
        ], counter=tc, caption="Prehliadka vnútorných rozvodov tepla a teplej vody")
        doc.add_paragraph("3.3 Zhodnotenie údržby a kontrola dokladov o údržbe a opravách rozvodov tepla a teplej vody").runs[0].bold = True
        _docx_table(doc, ["", ""], [
            ("Zhodnotenie údržby a zjavných stôp po údržbárskych prácach", data.get('roz_udrzba_stopy', '')),
            ("Kontrola dokladov o údržbe a opravách", data.get('roz_udrzba_doklady', '')),
        ], counter=tc, caption="Zhodnotenie údržby rozvodov tepla a teplej vody")
        doc.add_paragraph("3.4 Porovnanie skutočného využívania budovy s projektovaným využívaním").runs[0].bold = True
        _docx_table(doc, ["", ""], [
            ("Zhodnotenie skutočného využívania s projektovaným stavom využívania budovy", data.get('buduva_zhodnotenie', '')),
            ("Využívanie budovy od poslednej kontroly", data.get('zmena_vyuzivania', '')),
        ], counter=tc, caption="Porovnanie skutočného využívania budovy s projektovaným využívaním")
        doc.add_paragraph("3.5 Porovnanie skutočného využívania rozvodov tepla s projektovaným využívaním").runs[0].bold = True
        _docx_table(doc, ["", ""], [
            ("Zhodnotenie skutočného využívania s projektovaným stavom využívania rozvodov tepla", data.get('rozvody_zhodnotenie', '')),
            ("Využívanie rozvodov tepla od poslednej kontroly", data.get('rozvody_zmena', '')),
        ], counter=tc, caption="Porovnanie skutočného využívania rozvodov tepla s projektovaným využívaním")

        # --- 4. MERANIE ÚČINNOSTI KOTLA ---
        _docx_section_heading(doc, "4. Meranie účinnosti kotla – komínová strata")
        doc.add_paragraph("Účinnosť kontrolovaného kotla bola zistená nepriamou metódou z analýzy spalín.")
        if data.get('analyzator_typ') or data.get('analyzator_vc'):
            doc.add_paragraph(f"Meranie bolo vykonané analyzátorom spalín typ {data.get('analyzator_typ', '')}, {data.get('analyzator_vc', '')}. Pri kontrole účinnosti bol spaľovaný {data.get('druh_paliva', '')}.")
        for k in kotly:
            if k['merania']:
                header = [""] + [m['zatazenie'] for m in k['merania']]
                riadky = [
                    ["Teplota spaľovacieho vzduchu °C"] + [m['t_vzduch'] for m in k['merania']],
                    ["Teplota spalín °C"] + [m['t_spalin'] for m in k['merania']],
                    ["Obsah O2 v spalinách %"] + [m['o2'] for m in k['merania']],
                    ["Obsah CO v spalinách ppm"] + [m['co'] or '-' for m in k['merania']],
                    ["Obsah CO2 v spalinách %"] + [m['co2'] or '-' for m in k['merania']],
                    ["Obsah SO2 v spalinách ppm"] + [m['so2'] or '-' for m in k['merania']],
                    ["Obsah NO v spalinách ppm"] + [m['no'] or '-' for m in k['merania']],
                    ["Obsah NO2 v spalinách ppm"] + [m['no2'] or '-' for m in k['merania']],
                    ["Prebytok vzduchu -"] + [m['prebytok_vzduchu'] for m in k['merania']],
                    ["Strata kotla sálaním %"] + [m['strata_salanim'] for m in k['merania']],
                    ["Strata horľavinou v spalinách %"] + [m['strata_horlavinou'] for m in k['merania']],
                    ["Strata citeľným teplom spalín %"] + [m['strata_citelnym_teplom'] for m in k['merania']],
                    ["Účinnosť kotla %"] + [m['ucinnost'] for m in k['merania']],
                ]
                _docx_table(doc, header, riadky, counter=tc, caption=f"Namerané a vypočítané parametre – kotol {k['oznacenie']}")
            doc.add_paragraph(f"Priemerná účinnosť kotla: {k['priemerna_ucinnost']} %")
            if k['garantovana_ucinnost']:
                doc.add_paragraph(f"Garantovaná účinnosť kotla podľa výrobcu ȠG = {k['garantovana_ucinnost']}%")
            doc.add_paragraph(f"Minimálna požadovaná účinnosť kotla podľa Vyhl.č. 328/2005: ȠMIN = {vysledky.get('min_ucinnost')}%")
            p = doc.add_paragraph()
            p.add_run(f"Vyhodnotenie merania energetickej účinnosti: {k['stav']}").bold = True
            doc.add_paragraph("")

        # --- 5. PRIAMA METÓDA ---
        _docx_section_heading(doc, "5. Vyhodnotenie účinnosti výroby tepla priamou metódou")
        doc.add_paragraph("Spotrebu paliva pri výrobe tepla za predchádzajúce kalendárne roky uvádzam v nasledujúcej tabuľke.")
        if roky:
            header = [""] + [str(r['rok']) for r in roky]
            riadky = [
                ["Priemerné spaľovacie teplo objemové (kWh/jedn.)"] + [r['spalovacie_teplo'] or '-' for r in roky],
                ["Priemerná výhrevnosť paliva (kWh/jedn.)"] + [r['vyhrevnost'] for r in roky],
                ["Odber paliva"] + [r['odber_paliva'] for r in roky],
                ["Energia v palive / spotreba paliva (kWh)"] + [r['teplo_v_palive'] for r in roky],
                ["Pomer výhrevnosti k spaľovaciemu teplu (-)"] + [r['pomer_vyhrevnosti'] or '-' for r in roky],
                ["Teplo v palive (kWh)"] + [r['teplo_v_palive'] for r in roky],
                ["Vyrobené teplo (kWh)"] + [r['vyrobene_teplo'] for r in roky],
                ["Účinnosť výroby tepla (%)"] + [r['ucinnost_priama'] for r in roky],
                ["Spotreba tepla na vykurovanie (kWh)"] + [r['spotreba_vykurovanie'] for r in roky],
                ["Spotreba tepla na prípravu teplej vody (kWh)"] + [r['spotreba_tuv'] for r in roky],
                ["Spotreba vody na prípravu teplej vody (m³)"] + [data.get(f"rok{i+1}_spotreba_vody_tuv", '') for i in range(len(roky))],
                ["Dosiahnutá merná spotr. tepla na prípravu TV (kWh/m³)"] + [r['merna_spotreba_tuv'] for r in roky],
                ["Podiel spotreby teplej vody z celkovej spotreby (%)"] + [r['podiel_tuv'] for r in roky],
                ["Spotreba paliva na prípravu vykurovacej vody (kWh)"] + [r['spotreba_paliva_vykurovanie'] for r in roky],
                ["Spotreba paliva na prípravu teplej vody (kWh)"] + [r['spotreba_paliva_tuv'] for r in roky],
            ]
            _docx_table(doc, header, riadky, counter=tc, caption="Vyhodnotenie účinnosti výroby tepla priamou metódou")
            doc.add_paragraph(f"Priemerná účinnosť výroby tepla: {vysledky.get('priemerna_ucinnost_priama')} %")
        else:
            doc.add_paragraph("Neboli zadané ročné údaje o spotrebe paliva.")

        # --- 6. POSÚDENIE VÝKONU ---
        _docx_section_heading(doc, "6. Posúdenie výkonu kotla vzhľadom na potrebu tepla v budove")
        if roky:
            header = ["", "kWh/r", "GJ/r", "ti - tepr", "ti - te", "ti - tepr/ti - te", "kW"]
            riadky = [[r['rok'], r['spotreba_vykurovanie'], r['gj_rok'], round(r['ti'] - r['tepr'], 1),
                       round(r['ti'] - r['te'], 1), r['pomer_ti'], r['potrebny_vykon']] for r in roky]
            _docx_table(doc, header, riadky, counter=tc, caption="Zhodnotenie priemerného tepelného výkonu")
        doc.add_paragraph(f"Inštalovaný výkon kotlov v objekte spolu: {vysledky.get('instalovany_vykon')} kW")
        doc.add_paragraph(f"Potreba tepelného výkonu v objekte pre reálne podmienky: {vysledky.get('potrebny_vykon')} kW")
        rezerva = vysledky.get('vykon_rezerva', 0)
        if rezerva >= 0:
            doc.add_paragraph("Inštalovaný výkon kotolne pokrýva potrebu tepelného výkonu v objekte s dostatočnou rezervou.")
        else:
            doc.add_paragraph(f"Inštalovaný výkon kotolne nepokrýva potrebu tepelného výkonu v objekte (chýba {abs(rezerva)} kW).")

        # --- 7. VYHODNOTENIE KONTROLY KOTLA ---
        _docx_section_heading(doc, "7. Vyhodnotenie kontroly kotla a návrhy na opatrenia")
        doc.add_paragraph(f"A./ {data.get('navrh_a_text', '')}")
        p = doc.add_paragraph()
        p.add_run(f"B./ Kotol má mať v zmysle Vyhlášky č. 328/2005 Z.z. minimálnu priemernú účinnosť {vysledky.get('min_ucinnost')}%. Nameraná hodnota je:").bold = False
        for k in kotly:
            doc.add_paragraph(f"    {k['oznacenie']} = {k['priemerna_ucinnost']} %")
        doc.add_paragraph(f"Priemerná účinnosť výroby tepla zistená nepriamou metódou: {vysledky.get('priemerna_ucinnost')} %")
        doc.add_paragraph(f"Priemerná účinnosť výroby tepla zistená priamou metódou: {vysledky.get('priemerna_ucinnost_priama')} %")
        p = doc.add_paragraph()
        p.add_run(f"Kotly {'spĺňajú' if vysledky.get('celkovy_stav') == 'vyhovuje' else 'nespĺňajú'} požiadavky Vyhlášky č. 328/2005 Z.z.").bold = True

        # --- 8. VYHODNOTENIE ROZŠÍRENEJ KONTROLY ---
        _docx_section_heading(doc, "8. Vyhodnotenie rozšírenej kontroly vykurovacieho systému a návrh opatrení")
        doc.add_paragraph(f"C./ Rozšírenou kontrolou rozvodov vykurovania a teplej úžitkovej vody boli zistené tieto skutočnosti: {data.get('rozsirena_zistenia', '')}")
        doc.add_paragraph(f"D./ Pre zlepšenie prevádzkového stavu navrhujeme tieto opatrenia: {data.get('navrh_opatreni', '') or '—'}")
        doc.add_paragraph("Kontrola bola vykonaná podľa Zákona č. 314/2012 Z.z. § 3, v intervale podľa § 4, príloha č.1 o pravidelnej kontrole kotlov, vykurovacieho systému a klimatizačného systému, v rozsahu podľa Vyhlášky č. 422/2012 Z.z., ktorou sa ustanovuje postup pri pravidelnej kontrole vykurovacieho systému, rozšírenej kontrole vykurovacieho systému a pri pravidelnej kontrole klimatizačného systému.")
        doc.add_paragraph(f"Nasledujúcu kontrolu v zmysle Zákona č. 314/2012, príloha č. 1 je potrebné vykonať do: {vysledky.get('nasledujuca_kontrola')}")
        doc.add_paragraph("")
        doc.add_paragraph("Vlastník, prevádzkovateľ : .....................................")
        doc.add_paragraph("Oprávnená osoba : .....................................")
        doc.add_paragraph(f"Dňa : {data.get('datum_kontroly', '.....................................')}")
        doc.add_paragraph("")
        doc.add_paragraph(f"Vypracoval : {data.get('vypracoval_firma', '')}")

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    doc.add_paragraph(f"Dátum: {datetime.now().strftime('%d.%m.%Y')}")
    doc.add_paragraph(f"Budova: {data.get('nazov_budovy', '')}")
    doc.add_paragraph(f"Adresa: {data.get('adresa', '')}")
    doc.add_paragraph(f"Vlastník: {data.get('vlastnik', '')}")
    doc.add_paragraph("")

    doc.add_heading("Vstupné údaje", level=1)
    for k, v in data.items():
        if k not in ['nazov_budovy', 'adresa', 'vlastnik']:
            doc.add_paragraph(f"{k.replace('_', ' ').capitalize()}: {v}")

    doc.add_heading("Výsledky", level=1)
    for k, v in vysledky.items():
        doc.add_paragraph(f"{k.replace('_', ' ').capitalize()}: {v}")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

# ============================================================
# GENEROVANIE PDF - pomocné funkcie
# ============================================================
_pdf_cell_style = ParagraphStyle('pdfcell', fontName=PDF_FONT, fontSize=8.5, leading=11, alignment=TA_LEFT)
_pdf_cell_style_bold = ParagraphStyle('pdfcellb', fontName=PDF_FONT_BOLD, fontSize=8.5, leading=11, alignment=TA_LEFT)

def _pdf_table(rows, col_widths=None, header=True):
    """Tabuľka so zalamovaním textu v bunkách (Paragraph), pevnou šírkou 1. stĺpca
    a jemným zeleným podfarbením 1. stĺpca a hlavičky."""
    ncols = len(rows[0]) if rows else 1
    if col_widths is None:
        rest = (TABLE_TOTAL_CM - FIRST_COL_CM) / max(ncols - 1, 1)
        col_widths = [FIRST_COL_CM * cm] + [rest * cm] * (ncols - 1)
    else:
        col_widths = [FIRST_COL_CM * cm] + list(col_widths[1:])

    wrapped = []
    for ridx, row in enumerate(rows):
        wrapped_row = []
        for val in row:
            style = _pdf_cell_style_bold if (header and ridx == 0) else _pdf_cell_style
            text = '' if val in (None, '') else str(val)
            wrapped_row.append(Paragraph(text, style))
        wrapped.append(wrapped_row)

    t = Table(wrapped, colWidths=col_widths)
    t.hAlign = 'LEFT'
    style = [
        ('GRID', (0, 0), (-1, -1), 0.6, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BACKGROUND', (0, 0), (0, -1), TABLE_SHADE_RL),
    ]
    if header:
        style.append(('BACKGROUND', (0, 0), (-1, 0), TABLE_SHADE_RL))
    t.setStyle(TableStyle(style))
    return t

def _pdf_add_table(story, counter, caption, rows, col_widths=None, header=True):
    """Pridá do story číslovaný, vľavo zarovnaný popisok 'Tabuľka č. N' a tabuľku."""
    counter[0] += 1
    cap_style = ParagraphStyle('cap', fontName=PDF_FONT, fontSize=9, alignment=TA_LEFT,
                                italic=True, spaceAfter=3)
    label = f"Tabuľka č. {counter[0]} – {caption}" if caption else f"Tabuľka č. {counter[0]}"
    story.append(Paragraph(f"<i>{label}</i>", cap_style))
    story.append(_pdf_table(rows, col_widths=col_widths, header=header))
    story.append(Spacer(1, 0.4*cm))

def _pdf_header_footer(objekt, lokalita, firma):
    """Hlavička (Objekt, Lokalita + podčiarknutie) a päta (názov firmy) - okrem prvých 2 strán."""
    def _draw(canvas, doc):
        if doc.page <= 2:
            return
        canvas.saveState()
        canvas.setFont(PDF_FONT, 9)
        header_text = ", ".join([t for t in [objekt, lokalita] if t])
        header_y = A4[1] - 1.3*cm
        canvas.drawCentredString(A4[0] / 2, header_y, header_text)
        canvas.line(0, header_y - 0.15*cm, A4[0], header_y - 0.15*cm)
        canvas.setFont(PDF_FONT, 9)
        canvas.drawCentredString(A4[0] / 2, 1.3*cm, str(firma or ''))
        canvas.restoreState()
    return _draw

# ============================================================
# GENEROVANIE PDF
# ============================================================
def generuj_pdf(typ, data, vysledky):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2.2*cm)
    styles = getSampleStyleSheet()
    story = []

    nazov = {
        'audit': 'ENERGETICKÝ AUDIT',
        'certifikat': 'ENERGETICKÝ CERTIFIKÁT',
        'vykurovanie': 'Správa z vykonanej pravidelnej kontroly vykurovacieho systému'
    }.get(typ, 'SPRÁVA')

    if typ == 'vykurovanie':
        title_style = ParagraphStyle('title', parent=styles['Title'], fontName=PDF_FONT_BOLD,
                                      fontSize=16, leading=20, alignment=TA_CENTER,
                                      textColor=colors.HexColor('#8B1A1A'), spaceAfter=6)
        h1 = ParagraphStyle('h1', parent=styles['Heading1'], fontName=PDF_FONT_BOLD, fontSize=12.5,
                             textColor=colors.black, spaceBefore=14, spaceAfter=8, alignment=TA_LEFT)
        normal = ParagraphStyle('n', parent=styles['Normal'], fontName=PDF_FONT, fontSize=9.5, leading=13, alignment=TA_LEFT)
        bold = ParagraphStyle('b', parent=normal, fontName=PDF_FONT_BOLD)

        kotly = vysledky.get('kotly', [])
        roky = vysledky.get('roky', [])
        ma_oze = any(data.get(k) for k in ['oze_druh_energie', 'oze_druh_zariadenia', 'oze_vyrobca'])
        ma_tc = any(data.get(k) for k in ['tc_vyrobca', 'tc_typ', 'tc_prevadzkovy_stav'])
        tc = [0]  # zdieľané počítadlo tabuliek

        story.append(Paragraph(nazov, title_style))
        story.append(Paragraph("podľa Zákona č. 314/2012 Z.z.", ParagraphStyle('sub', parent=normal, alignment=TA_CENTER)))
        story.append(Spacer(1, 1*cm))

        for label, key in [
            ("Objekt", 'nazov_budovy'), ("Lokalita", 'adresa'), ("Majiteľ", 'vlastnik'),
            ("Adresa vlastníka", 'adresa_vlastnika'), ("Správca", 'spravca'), ("Prevádzkovateľ", 'prevadzkovatel'),
        ]:
            story.append(Paragraph(f"<b>{label}:</b> {data.get(key, '')}", normal))
        story.append(Spacer(1, 0.6*cm))

        fotografia = data.get('fotografia_cesta')
        if fotografia and os.path.isfile(fotografia):
            try:
                img = RLImage(fotografia, width=9*cm, height=6.5*cm, kind='proportional')
                img.hAlign = 'CENTER'
                story.append(img)
                story.append(Spacer(1, 0.6*cm))
            except Exception:
                pass

        for label, key in [
            ("Vypracoval", 'vypracoval_firma'), ("Kontakt", 'vypracoval_kontakt'),
            ("Dátum kontroly", 'datum_kontroly'), ("Poradové číslo", 'poradove_cislo'),
        ]:
            story.append(Paragraph(f"<b>{label}:</b> {data.get(key, '')}", normal))
        story.append(PageBreak())

        # --- OBSAH (strana 2 - bez hlavičky/päty) ---
        story.append(Paragraph("Obsah", h1))
        for polozka in [
            "1. Kontrola kotla",
            "&nbsp;&nbsp;&nbsp;&nbsp;1.1 Identifikačné údaje kontrolovaného kotla",
            "&nbsp;&nbsp;&nbsp;&nbsp;1.2 Identifikačné údaje ostatných zariadení na výrobu tepla",
            "2. Vizuálna kontrola a zhodnotenie funkčnosti a údržby kotla",
            "&nbsp;&nbsp;&nbsp;&nbsp;2.1 Prevádzková dokumentácia kotlov a povinnosti z nej vyplývajúce",
            "&nbsp;&nbsp;&nbsp;&nbsp;2.2 Vizuálna kontrola a zhodnotenie kotla",
            "&nbsp;&nbsp;&nbsp;&nbsp;2.3 Zhodnotenie údržby a kontrola dokladov o údržbe a opravách",
            "&nbsp;&nbsp;&nbsp;&nbsp;2.4 Kontrola funkčnosti kotla",
            "3. Rozšírená kontrola vykurovacieho systému",
            "&nbsp;&nbsp;&nbsp;&nbsp;3.1 Kontrola úplnosti a aktuálnosti dokumentácie vnútorných rozvodov tepla a teplej vody",
            "&nbsp;&nbsp;&nbsp;&nbsp;3.2 Prehliadka vnútorných rozvodov tepla a teplej vody",
            "&nbsp;&nbsp;&nbsp;&nbsp;3.3 Zhodnotenie údržby a kontrola dokladov o údržbe a opravách rozvodov tepla a teplej vody",
            "&nbsp;&nbsp;&nbsp;&nbsp;3.4 Porovnanie skutočného využívania budovy s projektovaným využívaním",
            "&nbsp;&nbsp;&nbsp;&nbsp;3.5 Porovnanie skutočného využívania rozvodov tepla s projektovaným využívaním",
            "4. Meranie účinnosti kotla – komínová strata",
            "5. Vyhodnotenie účinnosti výroby tepla priamou metódou",
            "6. Posúdenie výkonu kotla vzhľadom na potrebu tepla v budove",
            "7. Vyhodnotenie kontroly kotla a návrhy na opatrenia",
            "8. Vyhodnotenie rozšírenej kontroly vykurovacieho systému a návrh opatrení",
        ]:
            story.append(Paragraph(polozka, normal))

        story.append(Spacer(1, 0.4*cm))
        story.append(Paragraph("Zoznam tabuliek", bold))
        for i, caption in enumerate(_tabulky_zoznam(kotly, roky, ma_oze, ma_tc), start=1):
            story.append(Paragraph(f"Tabuľka č. {i} – {caption}", normal))
        story.append(PageBreak())

        # --- 1. KONTROLA KOTLA ---
        story.append(Paragraph("1. Kontrola kotla", h1))
        story.append(Paragraph("Kontrola kotla bola vykonávaná podľa Vyhl. č. 422/2012 Z.z. §2 v nasledovnom členení:", normal))
        story.append(Paragraph("1.1 Identifikačné údaje kontrolovaného kotla", bold))
        _pdf_add_table(story, tc, "Identifikačné údaje kontrolovaného kotla (spoločné údaje)", [
            ["Vlastník", data.get('vlastnik', '')], ["Adresa vlastníka", data.get('adresa_vlastnika', '')],
            ["Adresa budovy, v ktorej je kotol umiestnený", data.get('adresa', '')],
            ["Správca", data.get('spravca', '')], ["Prevádzkovateľ", data.get('prevadzkovatel', '')],
            ["Typ paliva", data.get('typ_paliva_kategoria', '')], ["Druh paliva", data.get('druh_paliva', '')],
            ["Spôsob dávkovania paliva", data.get('sposob_davkovania', '')],
        ], header=False)
        if kotly:
            header_row = [""] + [k['oznacenie'] for k in kotly]
            riadky = [
                ("Prevádzkový stav", 'prevadzkovy_stav'), ("Výrobca kotla", 'vyrobca'), ("Typ kotla", 'typ'),
                ("Výrobné číslo kotla", 'vyrobne_cislo'), ("Rok výroby kotla", 'rok_vyroby'),
                ("Menovitý výkon", 'menovity_vykon'), ("Max. výkon 50/30°C (kondenz.)", 'max_vykon_kondenzacny'),
                ("Max. výkon 80/60°C", 'max_vykon'), ("Max. príkon", 'max_prikon'),
                ("Min. výkon 80/60°C", 'min_vykon'), ("Min. príkon", 'min_prikon'),
                ("Kondenzačný/nekondenzačný", 'kondenzacny'), ("Označenie CE", 'oznacenie_ce'),
                ("Výrobca horáka", 'vyrobca_horaka'), ("Typ horáka", 'typ_horaka'),
                ("Výrobné číslo horáka", 'vyrobne_cislo_horaka'), ("Rok výroby horáka", 'rok_vyroby_horaka'),
            ]
            data_rows = [header_row] + [[label] + [str(k.get(key, '')) for k in kotly] for label, key in riadky]
            colN = (TABLE_TOTAL_CM - FIRST_COL_CM) / len(kotly)
            _pdf_add_table(story, tc, "Identifikačné údaje kontrolovaného kotla (parametre kotlov)",
                            data_rows, col_widths=[None] + [colN*cm]*len(kotly))
        else:
            story.append(Paragraph("Nebol zadaný žiadny kotol.", normal))
        _pdf_add_table(story, tc, "Identifikačné údaje kontrolovaného kotla (regulácia a médiá)", [
            ["Typ výkonovej regulácie", data.get('typ_regulacie', '')],
            ["Spôsob odvodu spalín", data.get('sposob_odvodu_spalin', '')],
            ["Spôsob prívodu vzduchu", data.get('sposob_privodu_vzduchu', '')],
            ["Teplonosné médium", data.get('teplonosne_medium', '')],
            ["Spôsob využitia kotla", data.get('sposob_vyuzitia', '')],
        ], header=False)

        story.append(Paragraph("1.2 Identifikačné údaje ostatných zariadení na výrobu tepla", bold))
        story.append(Paragraph("V budove je inštalované zariadenie na využívanie OZE." if ma_oze else "V budove nie je inštalované zariadenie na využívanie OZE.", normal))
        if ma_oze:
            _pdf_add_table(story, tc, "Identifikačné údaje ostatných zariadení na výrobu tepla – OZE", [
                ["Druh využívanej energie", data.get('oze_druh_energie', '')],
                ["Druh zariadenia", data.get('oze_druh_zariadenia', '')], ["Výrobca", data.get('oze_vyrobca', '')],
                ["Apertúrna plocha kolektora", data.get('oze_apertura', '')], ["Počet kusov", data.get('oze_pocet_kusov', '')],
                ["Celkový inštalovaný výkon", data.get('oze_celkovy_vykon', '')],
            ], header=False)
        if ma_tc:
            story.append(Paragraph("Tepelné čerpadlo", normal))
            _pdf_add_table(story, tc, "Identifikačné údaje ostatných zariadení na výrobu tepla – tepelné čerpadlo", [
                ["Prevádzkový stav", data.get('tc_prevadzkovy_stav', '')], ["Výrobca", data.get('tc_vyrobca', '')],
                ["Typ", data.get('tc_typ', '')], ["Menovitý výkon", data.get('tc_menovity_vykon', '')],
                ["Menovitý príkon", data.get('tc_menovity_prikon', '')], ["COP", data.get('tc_cop', '')],
            ], header=False)
        if data.get('poznamky_zariadenia'):
            story.append(Paragraph(f"Poznámky: {data.get('poznamky_zariadenia')}", normal))
        story.append(PageBreak())

        # --- 2. VIZUÁLNA KONTROLA ---
        story.append(Paragraph("2. Vizuálna kontrola a zhodnotenie funkčnosti a údržby kotla", h1))
        story.append(Paragraph("2.1 Prevádzková dokumentácia kotlov a povinnosti z nej vyplývajúce", bold))
        _pdf_add_table(story, tc, "Prevádzková dokumentácia kotlov", [["Dokumentácia kotla", ""]] + [
            [label, data.get(key, '') or '—'] for label, key in [
                ("Projektová dokumentácia kotla", 'dok_projektova'), ("Prevádzkový predpis výr. kotla", 'dok_predpis'),
                ("Dokumentácia prevádzky a údržby", 'dok_udrzba'), ("Správa z predchádzajúcej kontroly", 'dok_predchadzajuca_sprava'),
                ("Zaškolenie obsluhy", 'dok_zaskolenie'), ("Odborné prehliadky a revízie", 'dok_revizie'),
            ]
        ])
        story.append(Paragraph("2.2 Vizuálna kontrola a zhodnotenie kotla", bold))
        _pdf_add_table(story, tc, "Vizuálna kontrola a zhodnotenie kotla", [["", ""]] + [
            [label, data.get(key, '') or '—'] for label, key in [
                ("Únik paliva", 'unik_paliva'), ("Únik teplonosnej látky", 'unik_teplonosnej_latky'),
                ("Vonkajší stav kotla", 'vonkajsi_stav_kotla'),
                ("Tepelná izolácia, oplechovanie, netesnosti spalinovodu", 'izolacia_oplechovanie'),
                ("Znečistenie spaľovacej komory a teplovýmenných plôch", 'znecistenie'),
                ("Znečistenie horákov", 'znecistenie_horakov'),
                ("Funkčnosť armatúr a ostatných častí", 'armatury'),
                ("Kvalita teplonosnej látky, čistota obehovej vody", 'kvalita_vody'),
                ("Správnosť údajov meracích prístrojov", 'meracie_pristroje'),
                ("Systém riadenia kotla", 'system_riadenia'), ("Čistota a poriadok kotolne", 'cistota_kotolne'),
            ]
        ])
        story.append(PageBreak())

        story.append(Paragraph("2.3 Zhodnotenie údržby a kontrola dokladov o údržbe a opravách", bold))
        _pdf_add_table(story, tc, "Zhodnotenie údržby a kontrola dokladov o údržbe a opravách kotla", [["", ""]] + [
            [label, data.get(key, '') or '—'] for label, key in [
                ("Zjavné stopy po údržbárskych prácach", 'udrzba_stopy'), ("Doklady o údržbe a opravách", 'doklady_udrzba'),
            ]
        ])
        story.append(Paragraph("2.4 Kontrola funkčnosti kotla", bold))
        funkcnost_rows = [["", ""], ["Skúška funkcie kotlov v prevádzke", data.get('funkcnost_skuska', '') or '—']]
        for k in kotly:
            if k['overenie_min'] or k['overenie_max']:
                funkcnost_rows.append([f"Overenie výkonu kotla {k['oznacenie']}", f"min = {k['overenie_min']} m³/hod   max = {k['overenie_max']} m³/hod"])
        funkcnost_rows.append(["Regulácia výkonu", data.get('regulacia_vykonu_text', '') or '—'])
        _pdf_add_table(story, tc, "Kontrola funkčnosti kotla", funkcnost_rows)

        # --- 3. ROZŠÍRENÁ KONTROLA ---
        story.append(Paragraph("3. Rozšírená kontrola vykurovacieho systému", h1))
        story.append(Paragraph("Rozšírená kontrola vykurovacieho systému bola vykonávaná podľa Vyhl. č. 422/2012 Z.z. §3 v nasledovnom členení:", normal))
        story.append(Paragraph("3.1 Kontrola úplnosti a aktuálnosti dokumentácie vnútorných rozvodov tepla a teplej vody", bold))
        _pdf_add_table(story, tc, "Kontrola úplnosti a aktuálnosti dokumentácie vnútorných rozvodov", [["", ""]] + [
            [label, data.get(key, '') or '—'] for label, key in [
                ("Úplnosť PD ÚK a TÚV", 'roz_pd_uk_tuv'), ("Úplnosť PD zmien a rekonštrukcie", 'roz_pd_zmien'),
                ("Úplnosť prevádz. predpisov zariadení", 'roz_predpisy'), ("Úplnosť MPP", 'roz_mpp'),
                ("Vedenie prevádzkového denníka", 'roz_dennik'), ("Správy o údržbe a opravách", 'roz_spravy_udrzba'),
                ("Správa z predchádzajúcej kontroly", 'dok_predchadzajuca_sprava'),
            ]
        ])
        story.append(PageBreak())

        story.append(Paragraph("3.2 Prehliadka vnútorných rozvodov tepla a teplej vody", bold))
        _pdf_add_table(story, tc, "Prehliadka vnútorných rozvodov tepla a teplej vody", [["", ""]] + [
            [label, data.get(key, '') or '—'] for label, key in [
                ("Hlavné komponenty rozv. tepla, prvky merania a riadenia", 'roz_komponenty'),
                ("Ovládacie prvky systému regulácie", 'roz_ovladacie_prvky'),
                ("Vykurovacie telesá", 'vykurovacie_telesa'), ("Tepelná izolácia rozvodov tepla", 'tepelna_izolacia_rozvodov'),
                ("Čistota obehovej vody", 'cistota_obehovej_vody'),
            ]
        ])
        story.append(Paragraph("3.3 Zhodnotenie údržby a kontrola dokladov o údržbe a opravách rozvodov tepla a teplej vody", bold))
        _pdf_add_table(story, tc, "Zhodnotenie údržby rozvodov tepla a teplej vody", [["", ""]] + [
            [label, data.get(key, '') or '—'] for label, key in [
                ("Zjavné stopy po údržbárskych prácach", 'roz_udrzba_stopy'), ("Doklady o údržbe a opravách", 'roz_udrzba_doklady'),
            ]
        ])
        story.append(Paragraph("3.4 Porovnanie skutočného využívania budovy s projektovaným využívaním", bold))
        _pdf_add_table(story, tc, "Porovnanie skutočného využívania budovy s projektovaným využívaním", [["", ""]] + [
            [label, data.get(key, '') or '—'] for label, key in [
                ("Zhodnotenie skutočného využívania s projektovaným stavom využívania budovy", 'buduva_zhodnotenie'),
                ("Využívanie budovy od poslednej kontroly", 'zmena_vyuzivania'),
            ]
        ])
        story.append(Paragraph("3.5 Porovnanie skutočného využívania rozvodov tepla s projektovaným využívaním", bold))
        _pdf_add_table(story, tc, "Porovnanie skutočného využívania rozvodov tepla s projektovaným využívaním", [["", ""]] + [
            [label, data.get(key, '') or '—'] for label, key in [
                ("Zhodnotenie skutočného využívania s projektovaným stavom využívania rozvodov tepla", 'rozvody_zhodnotenie'),
                ("Využívanie rozvodov tepla od poslednej kontroly", 'rozvody_zmena'),
            ]
        ])
        story.append(PageBreak())

        # --- 4. MERANIE ÚČINNOSTI KOTLA ---
        story.append(Paragraph("4. Meranie účinnosti kotla – komínová strata", h1))
        story.append(Paragraph("Účinnosť kontrolovaného kotla bola zistená nepriamou metódou z analýzy spalín.", normal))
        if data.get('analyzator_typ') or data.get('analyzator_vc'):
            story.append(Paragraph(f"Meranie bolo vykonané analyzátorom spalín typ {data.get('analyzator_typ', '')}, {data.get('analyzator_vc', '')}. Pri kontrole účinnosti bol spaľovaný {data.get('druh_paliva', '')}.", normal))
        for k in kotly:
            if k['merania']:
                header_row = [""] + [m['zatazenie'] for m in k['merania']]
                riadky = [
                    ["Teplota spaľovacieho vzduchu °C"] + [str(m['t_vzduch']) for m in k['merania']],
                    ["Teplota spalín °C"] + [str(m['t_spalin']) for m in k['merania']],
                    ["Obsah O2 v spalinách %"] + [str(m['o2']) for m in k['merania']],
                    ["Obsah CO v spalinách ppm"] + [str(m['co'] or '-') for m in k['merania']],
                    ["Obsah CO2 v spalinách %"] + [str(m['co2'] or '-') for m in k['merania']],
                    ["Obsah SO2 v spalinách ppm"] + [str(m['so2'] or '-') for m in k['merania']],
                    ["Obsah NO v spalinách ppm"] + [str(m['no'] or '-') for m in k['merania']],
                    ["Obsah NO2 v spalinách ppm"] + [str(m['no2'] or '-') for m in k['merania']],
                    ["Prebytok vzduchu -"] + [str(m['prebytok_vzduchu']) for m in k['merania']],
                    ["Strata kotla sálaním %"] + [str(m['strata_salanim']) for m in k['merania']],
                    ["Strata horľavinou v spalinách %"] + [str(m['strata_horlavinou']) for m in k['merania']],
                    ["Strata citeľným teplom spalín %"] + [str(m['strata_citelnym_teplom']) for m in k['merania']],
                    ["Účinnosť kotla %"] + [str(m['ucinnost']) for m in k['merania']],
                ]
                colN = (TABLE_TOTAL_CM - FIRST_COL_CM) / max(len(k['merania']), 1)
                _pdf_add_table(story, tc, f"Namerané a vypočítané parametre – kotol {k['oznacenie']}",
                                [header_row] + riadky, col_widths=[None] + [colN*cm]*len(k['merania']))
            story.append(Paragraph(f"Priemerná účinnosť kotla: <b>{k['priemerna_ucinnost']} %</b>", normal))
            if k['garantovana_ucinnost']:
                story.append(Paragraph(f"Garantovaná účinnosť kotla podľa výrobcu ȠG = {k['garantovana_ucinnost']}%", normal))
            story.append(Paragraph(f"Minimálna požadovaná účinnosť kotla podľa Vyhl.č. 328/2005: ȠMIN = {vysledky.get('min_ucinnost')}%", normal))
            story.append(Paragraph(f"<b>Vyhodnotenie merania energetickej účinnosti: {k['stav']}</b>", normal))
            story.append(Spacer(1, 0.4*cm))
        story.append(PageBreak())

        # --- 5. PRIAMA METÓDA ---
        story.append(Paragraph("5. Vyhodnotenie účinnosti výroby tepla priamou metódou", h1))
        story.append(Paragraph("Spotrebu paliva pri výrobe tepla za predchádzajúce kalendárne roky uvádzam v nasledujúcej tabuľke.", normal))
        if roky:
            header_row = [""] + [str(r['rok']) for r in roky]
            riadky = [
                ["Priemerné spaľovacie teplo objemové (kWh/jedn.)"] + [str(r['spalovacie_teplo'] or '-') for r in roky],
                ["Priemerná výhrevnosť paliva (kWh/jedn.)"] + [str(r['vyhrevnost']) for r in roky],
                ["Odber paliva"] + [str(r['odber_paliva']) for r in roky],
                ["Energia v palive (kWh)"] + [str(r['teplo_v_palive']) for r in roky],
                ["Pomer výhrevnosti k spaľovaciemu teplu (-)"] + [str(r['pomer_vyhrevnosti'] or '-') for r in roky],
                ["Vyrobené teplo (kWh)"] + [str(r['vyrobene_teplo']) for r in roky],
                ["Účinnosť výroby tepla (%)"] + [str(r['ucinnost_priama']) for r in roky],
                ["Spotreba tepla na vykurovanie (kWh)"] + [str(r['spotreba_vykurovanie']) for r in roky],
                ["Spotreba tepla na prípravu TÚV (kWh)"] + [str(r['spotreba_tuv']) for r in roky],
                ["Dosiahnutá merná spotr. tepla na TÚV (kWh/m³)"] + [str(r['merna_spotreba_tuv']) for r in roky],
                ["Podiel spotreby TÚV z celkovej spotreby (%)"] + [str(r['podiel_tuv']) for r in roky],
                ["Spotreba paliva na prípravu vykurovacej vody (kWh)"] + [str(r['spotreba_paliva_vykurovanie']) for r in roky],
                ["Spotreba paliva na prípravu teplej vody (kWh)"] + [str(r['spotreba_paliva_tuv']) for r in roky],
            ]
            colN = (TABLE_TOTAL_CM - FIRST_COL_CM) / len(roky)
            _pdf_add_table(story, tc, "Vyhodnotenie účinnosti výroby tepla priamou metódou",
                            [header_row] + riadky, col_widths=[None] + [colN*cm]*len(roky))
            story.append(Paragraph(f"Priemerná účinnosť výroby tepla: <b>{vysledky.get('priemerna_ucinnost_priama')} %</b>", normal))
        else:
            story.append(Paragraph("Neboli zadané ročné údaje o spotrebe paliva.", normal))
        story.append(PageBreak())

        # --- 6. POSÚDENIE VÝKONU ---
        story.append(Paragraph("6. Posúdenie výkonu kotla vzhľadom na potrebu tepla v budove", h1))
        if roky:
            header_row = ["", "kWh/r", "GJ/r", "ti - tepr", "ti - te", "pomer", "kW"]
            riadky = [[str(r['rok']), str(r['spotreba_vykurovanie']), str(r['gj_rok']), str(round(r['ti'] - r['tepr'], 1)),
                       str(round(r['ti'] - r['te'], 1)), str(r['pomer_ti']), str(r['potrebny_vykon'])] for r in roky]
            colN = (TABLE_TOTAL_CM - FIRST_COL_CM) / 6
            _pdf_add_table(story, tc, "Zhodnotenie priemerného tepelného výkonu",
                            [header_row] + riadky, col_widths=[None] + [colN*cm]*6)
        story.append(Paragraph(f"Inštalovaný výkon kotlov v objekte spolu: <b>{vysledky.get('instalovany_vykon')} kW</b>", normal))
        story.append(Paragraph(f"Potreba tepelného výkonu v objekte pre reálne podmienky: <b>{vysledky.get('potrebny_vykon')} kW</b>", normal))
        rezerva = vysledky.get('vykon_rezerva', 0)
        if rezerva >= 0:
            story.append(Paragraph("Inštalovaný výkon kotolne pokrýva potrebu tepelného výkonu v objekte s dostatočnou rezervou.", normal))
        else:
            story.append(Paragraph(f"Inštalovaný výkon kotolne nepokrýva potrebu tepelného výkonu v objekte (chýba {abs(rezerva)} kW).", normal))
        story.append(PageBreak())

        # --- 7. VYHODNOTENIE KONTROLY KOTLA ---
        story.append(Paragraph("7. Vyhodnotenie kontroly kotla a návrhy na opatrenia", h1))
        story.append(Paragraph(f"A./ {data.get('navrh_a_text', '')}", normal))
        story.append(Paragraph(f"B./ Kotol má mať v zmysle Vyhlášky č. 328/2005 Z.z. minimálnu priemernú účinnosť {vysledky.get('min_ucinnost')}%. Nameraná hodnota je:", normal))
        for k in kotly:
            story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;&nbsp;{k['oznacenie']} = {k['priemerna_ucinnost']} %", normal))
        story.append(Paragraph(f"Priemerná účinnosť výroby tepla zistená nepriamou metódou: {vysledky.get('priemerna_ucinnost')} %", normal))
        story.append(Paragraph(f"Priemerná účinnosť výroby tepla zistená priamou metódou: {vysledky.get('priemerna_ucinnost_priama')} %", normal))
        splna = 'spĺňajú' if vysledky.get('celkovy_stav') == 'vyhovuje' else 'nespĺňajú'
        story.append(Paragraph(f"<b>Kotly {splna} požiadavky Vyhlášky č. 328/2005 Z.z.</b>", normal))
        story.append(Spacer(1, 0.5*cm))

        # --- 8. VYHODNOTENIE ROZŠÍRENEJ KONTROLY ---
        story.append(Paragraph("8. Vyhodnotenie rozšírenej kontroly vykurovacieho systému a návrh opatrení", h1))
        story.append(Paragraph(f"C./ Rozšírenou kontrolou rozvodov vykurovania a teplej úžitkovej vody boli zistené tieto skutočnosti: {data.get('rozsirena_zistenia', '')}", normal))
        story.append(Paragraph(f"D./ Pre zlepšenie prevádzkového stavu navrhujeme tieto opatrenia: {data.get('navrh_opatreni', '') or '—'}", normal))
        story.append(Paragraph("Kontrola bola vykonaná podľa Zákona č. 314/2012 Z.z. § 3, v intervale podľa § 4, príloha č.1 o pravidelnej kontrole kotlov, vykurovacieho systému a klimatizačného systému, v rozsahu podľa Vyhlášky č. 422/2012 Z.z.", normal))
        story.append(Paragraph(f"Nasledujúcu kontrolu v zmysle Zákona č. 314/2012, príloha č. 1 je potrebné vykonať do: <b>{vysledky.get('nasledujuca_kontrola')}</b>", normal))
        story.append(Spacer(1, 1*cm))
        story.append(Paragraph("Vlastník, prevádzkovateľ : .....................................", normal))
        story.append(Paragraph("Oprávnená osoba : .....................................", normal))
        story.append(Paragraph(f"Dňa : {data.get('datum_kontroly', '.....................................')}", normal))
        story.append(Spacer(1, 0.5*cm))
        story.append(Paragraph(f"Vypracoval : {data.get('vypracoval_firma', '')}", normal))

        header_footer_fn = _pdf_header_footer(data.get('nazov_budovy'), data.get('adresa'), data.get('vypracoval_firma'))
        doc.build(story, onFirstPage=header_footer_fn, onLaterPages=header_footer_fn)
        buf.seek(0)
        return buf

    title_style = ParagraphStyle('title', parent=styles['Title'], fontSize=16, spaceAfter=20)
    story.append(Paragraph(nazov, title_style))
    story.append(Spacer(1, 0.5*cm))

    # Základné info
    info = [
        ['Dátum:', datetime.now().strftime('%d.%m.%Y')],
        ['Budova:', data.get('nazov_budovy', '')],
        ['Adresa:', data.get('adresa', '')],
        ['Vlastník:', data.get('vlastnik', '')],
    ]
    t = Table(info, colWidths=[5*cm, 12*cm])
    t.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.5*cm))

    story.append(Paragraph("Vstupné údaje", styles['Heading2']))
    for k, v in data.items():
        if k not in ['nazov_budovy', 'adresa', 'vlastnik']:
            story.append(Paragraph(f"<b>{k.replace('_',' ').capitalize()}:</b> {v}", styles['Normal']))

    story.append(Spacer(1, 0.5*cm))
    story.append(Paragraph("Výsledky", styles['Heading2']))

    vysl_data = [['Ukazovateľ', 'Hodnota']]
    for k, v in vysledky.items():
        vysl_data.append([k.replace('_',' ').capitalize(), str(v)])

    tv = Table(vysl_data, colWidths=[9*cm, 8*cm])
    tv.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2c7a4b')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f0f7f3')]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('TOPPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(tv)

    doc.build(story)
    buf.seek(0)
    return buf

# ============================================================
# DEKORATOR - prihlasenie
# ============================================================
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'logged_in' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

# ============================================================
# ROUTES
# ============================================================
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        if request.form['username'] == USERNAME and request.form['password'] == PASSWORD:
            session['logged_in'] = True
            return redirect(url_for('dashboard'))
        flash('Nesprávne meno alebo heslo.')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('index'))

@app.route('/dashboard')
@login_required
def dashboard():
    conn = get_db()
    zaznamy = conn.execute('SELECT * FROM audity ORDER BY datum DESC LIMIT 10').fetchall()
    conn.close()
    return render_template('dashboard.html', zaznamy=zaznamy)

# --- AUDIT ---
@app.route('/audit', methods=['GET', 'POST'])
@login_required
def audit():
    if request.method == 'POST':
        data = request.form.to_dict()
        vysledky = vypocitaj_audit(data)
        conn = get_db()
        conn.execute('INSERT INTO audity (typ, datum, nazov_budovy, adresa, vlastnik, udaje, vysledky) VALUES (?,?,?,?,?,?,?)',
            ('audit', datetime.now().strftime('%d.%m.%Y %H:%M'),
             data.get('nazov_budovy'), data.get('adresa'), data.get('vlastnik'),
             str(data), str(vysledky)))
        conn.commit()
        conn.close()
        return render_template('vysledok.html', typ='audit', data=data, vysledky=vysledky)
    return render_template('audit.html')

# --- CERTIFIKAT ---
@app.route('/certifikat', methods=['GET', 'POST'])
@login_required
def certifikat():
    if request.method == 'POST':
        data = request.form.to_dict()
        vysledky = vypocitaj_certifikat(data)
        conn = get_db()
        conn.execute('INSERT INTO audity (typ, datum, nazov_budovy, adresa, vlastnik, udaje, vysledky) VALUES (?,?,?,?,?,?,?)',
            ('certifikat', datetime.now().strftime('%d.%m.%Y %H:%M'),
             data.get('nazov_budovy'), data.get('adresa'), data.get('vlastnik'),
             str(data), str(vysledky)))
        conn.commit()
        conn.close()
        return render_template('vysledok.html', typ='certifikat', data=data, vysledky=vysledky)
    return render_template('certifikat.html')

# --- VYKUROVANIE ---
UPLOAD_DIR = os.path.join('static', 'uploads')

def uloz_fotografiu(subor):
    if not subor or not subor.filename:
        return ''
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    prípona = os.path.splitext(subor.filename)[1].lower()
    if prípona not in ('.jpg', '.jpeg', '.png'):
        return ''
    nazov = f"kotolna_{datetime.now().strftime('%Y%m%d%H%M%S%f')}{prípona}"
    subor.save(os.path.join(UPLOAD_DIR, nazov))
    return os.path.join(UPLOAD_DIR, nazov)

@app.route('/vykurovanie', methods=['GET', 'POST'])
@login_required
def vykurovanie():
    if request.method == 'POST':
        data = request.form.to_dict()
        data['fotografia_cesta'] = uloz_fotografiu(request.files.get('fotografia'))
        vysledky = vypocitaj_vykurovanie(data)
        conn = get_db()
        conn.execute('INSERT INTO audity (typ, datum, nazov_budovy, adresa, vlastnik, udaje, vysledky) VALUES (?,?,?,?,?,?,?)',
            ('vykurovanie', datetime.now().strftime('%d.%m.%Y %H:%M'),
             data.get('nazov_budovy'), data.get('adresa'), data.get('vlastnik'),
             str(data), str(vysledky)))
        conn.commit()
        conn.close()
        return render_template('vysledok.html', typ='vykurovanie', data=data, vysledky=vysledky)
    return render_template('vykurovanie.html')

# --- STIAHNUTIE ---
@app.route('/stiahnut/<format>/<typ>', methods=['POST'])
@login_required
def stiahnut(format, typ):
    data = request.form.to_dict()
    vysledky_fn = {'audit': vypocitaj_audit, 'certifikat': vypocitaj_certifikat, 'vykurovanie': vypocitaj_vykurovanie}
    vysledky = vysledky_fn[typ](data)

    nazov = data.get('nazov_budovy', 'sprava').replace(' ', '_')
    datum = datetime.now().strftime('%Y%m%d')

    if format == 'word':
        buf = generuj_word(typ, data, vysledky)
        return send_file(buf, as_attachment=True,
                        download_name=f"{typ}_{nazov}_{datum}.docx",
                        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    else:
        buf = generuj_pdf(typ, data, vysledky)
        return send_file(buf, as_attachment=True,
                        download_name=f"{typ}_{nazov}_{datum}.pdf",
                        mimetype='application/pdf')

@app.route('/historia')
@login_required
def historia():
    conn = get_db()
    zaznamy = conn.execute('SELECT * FROM audity ORDER BY datum DESC').fetchall()
    conn.close()
    return render_template('historia.html', zaznamy=zaznamy)

with app.app_context():
    init_db()

if __name__ == '__main__':
    app.run(debug=True)
